"""Build a blind human-annotation workbook for 10 of the 30 gold advertisements.
Run from step3_baseline_nlp/.  Writes human_annotation_sheet.xlsx (give to the annotator)
and human_annotation_key_PRIVATE.csv (keep away from the annotator)."""
import random, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from gold_annotations import GOLD
from skill_extractor_v2 import RequirementAwareExtractor
SEED=36118; N_ADS=10; N_DISTRACT=12
jobs=pd.read_csv("job_cleaned.csv",encoding="utf-8-sig")
BASE=set(pd.read_csv("skills_dictionary.csv").canonical_skill)
EXTRA=pd.read_csv("extra_aliases_reviewed.csv"); UNIV=BASE|set(EXTRA.canonical_skill)
rng=random.Random(SEED)
rows=sorted(GOLD)
# prefer ads that are not very long, so ten ads take about 30-40 minutes
rows_ok=[r for r in rows if len(jobs.job_description[r])<=7000]
pick=sorted(rng.sample(rows_ok,N_ADS))
ext=RequirementAwareExtractor(extra_aliases_path="extra_aliases_reviewed.csv")
wb=Workbook(); ws0=wb.active; ws0.title="Instructions"
instr=["HUMAN ANNOTATION TASK (about 30-40 minutes)",
"",
"You are the second, independent annotator. Please do NOT look at the project code, report or any tool output before finishing.",
"1. Open sheet 'Ads' and read advertisement 1. Then go to sheet 'Labels' and, for each skill listed for that ad, choose one label in column D:",
"   R = the ad clearly requires the skill or says the role will use it (requirements, responsibilities, 'you will...'). The skill is named in the text.",
"   P = the skill is named but only as preferred / nice-to-have / bonus.",
"   I = the skill is NOT named in the ad, but the ad clearly demands the underlying concept (for example 'measure how well our models perform' implies Model Evaluation).",
"   N = not demanded. Includes skills that appear only in the company description, benefits or legal text, and skills that are simply not in the ad.",
"2. Label every row; do not leave blanks. Use the ad text only, not outside knowledge of the company.",
"3. Skills appear in alphabetical order and many of them will be N. That is expected.",
"4. When finished, save the file and send it back unchanged in structure (keep the sheet names and columns).",
"Do not discuss your labels with anyone until you have sent the file back."]
for i,t in enumerate(instr,1): ws0.cell(i,1,t)
ws0["A1"].font=Font(bold=True,size=14); ws0.column_dimensions["A"].width=140
ws1=wb.create_sheet("Ads"); ws1.append(["ad_no","job_title","company","advertisement_text"])
ws2=wb.create_sheet("Labels"); ws2.append(["ad_no","job_title","skill","label (R / P / I / N)"])
key=[]
for no,r in enumerate(pick,1):
    text=str(jobs.job_description[r]); title=jobs.job_title[r]; comp=jobs.company[r]
    ws1.append([no,title,comp,text[:32000]])
    hits=set(ext.extract_skills(text)) & UNIV
    g=GOLD[r]; goldset=(g["R"]|g["P"]|g["I"])&UNIV
    pool=hits|goldset
    others=sorted(UNIV-pool); rng.shuffle(others); pool|=set(others[:N_DISTRACT])
    for s in sorted(pool, key=str.lower):
        ws2.append([no,title,s,None])
        lab="R" if s in g["R"] else "P" if s in g["P"] else "I" if s in g["I"] else "N"
        key.append(dict(ad_no=no,row_index=r,job_id=jobs.job_id[r],skill=s,gold_label=lab))
for c in ws1["D"][1:]: c.alignment=Alignment(wrap_text=True,vertical="top")
ws1.column_dimensions["B"].width=34; ws1.column_dimensions["C"].width=22; ws1.column_dimensions["D"].width=140
ws2.column_dimensions["B"].width=34; ws2.column_dimensions["C"].width=30; ws2.column_dimensions["D"].width=22
dv=DataValidation(type="list",formula1='"R,P,I,N"',allow_blank=True); ws2.add_data_validation(dv); dv.add(f"D2:D{ws2.max_row}")
for ws in (ws1,ws2):
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F3864")
    ws.freeze_panes="A2"
wb.save("human_annotation_sheet.xlsx")
pd.DataFrame(key).to_csv("human_annotation_key_PRIVATE.csv",index=False)
print(len(pick),"ads,",len(key),"labels; ads:",pick)
print(pd.DataFrame(key).gold_label.value_counts().to_dict(), "max text len",max(len(str(jobs.job_description[r])) for r in pick))
